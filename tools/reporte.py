"""Reportes diarios de vencimientos por gerente (Excel para OneDrive).

Uso:  python3 reporte.py <ultimo.json o '-'> <aplicativo index.html> <destinatarios.json> <carpeta_salida>
- ultimo.json: cierre que envía la hoja de Google a GitHub ('-' = sin datos, genera los archivos vacíos).
- aplicativo: página del aplicativo (sucursales, unidad, zona y SKU con inventario por tienda).
- destinatarios.json: [{"correo", "nombre", "archivo", "alcance": {"unidad": ...} | {"zona": ...} | {"todas": true}}]
Imprime JSON con el resumen por destinatario (lotes, unidades críticas) para armar los correos.
"""
import sys, json, datetime as dt, os, re, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule, FormulaRule

src, app, dest, out = sys.argv[1:5]
os.makedirs(out, exist_ok=True)
page = open(app, encoding='utf-8').read()
D = json.JSONDecoder().raw_decode(page[page.index('const D = ') + 10:])[0]
datos = json.load(open(src, encoding='utf-8')) if src != '-' else {'cierre': (dt.date.today() - dt.timedelta(days=1)).isoformat(), 'filas': [], 'avance': {}}
cierre = dt.date.fromisoformat(datos['cierre']); hoy = cierre + dt.timedelta(days=1)
MESES = ['enero','febrero','marzo','abril','mayo','junio','julio','agosto','septiembre','octubre','noviembre','diciembre']
mes = f'{cierre.year}-M{cierre.month:02d}'; mes_txt = MESES[cierre.month - 1]   # periodo mensual del aplicativo (texto, ej. 2026-M10)
corto = lambda s: re.sub(r'^[^|]+\| ', '', s)

# Último reporte por sucursal + código (todas sus filas con la misma marca de tiempo)
grupos = collections.defaultdict(list)
for f in datos['filas']: grupos[(f[3], f[5])].append(f)
# Inventario en sistema de hoy por sucursal y código (viene en el aplicativo, se actualiza de lunes a viernes)
inv_hoy = {}
for si, s in enumerate(D['s']):
  a = D['i'][si]; inv_hoy[s[2]] = {D['p'][a[k]][0]: a[k + 1] for k in range(0, len(a), 2)}
lotes = []
for (suc, cod), fs in grupos.items():
  ult = max(f[13] for f in fs)
  fs = sorted([f for f in fs if f[13] == ult], key=lambda f: f[10])          # lotes del último conteo, del más próximo a vencer al más lejano
  inv_conteo = fs[0][7]; hoy_u = inv_hoy.get(suc, {}).get(cod, 0)
  salio = max(0, inv_conteo - hoy_u)                                         # unidades que salieron desde el conteo (venta, merma, traslado)
  quedan_max = hoy_u
  for f in fs:
    v = dt.date.fromisoformat(f[10][:10]); dias = (v - hoy).days
    ajust = max(0, f[9] - salio); salio = max(0, salio - f[9])               # se asume que sale primero lo que vence primero
    ajust = min(ajust, quedan_max); quedan_max -= ajust                       # nunca más de lo que hay hoy en sistema
    reg = dt.datetime.fromisoformat(f[13]); edad = (hoy - reg.date()).days
    lotes.append({'suc': suc, 'marca': f[4], 'cod': cod, 'prod': f[6], 'inv': inv_conteo, 'inv_hoy': hoy_u, 'rep': f[9], 'u': ajust, 'v': v, 'dias': dias,
                  'col': f[11], 'reg': reg, 'edad': edad, 'nuevo': f[13][:10] == cierre.isoformat()})
def accion(l):
  if l['u'] == 0: return 'Ya no hay unidades (vendido o retirado)'
  if l['dias'] < 0: return 'Retirar del anaquel y registrar merma'
  if l['dias'] <= 30 and l['edad'] > 7: return 'Recontar (crítico, conteo de hace más de 7 días)'
  if l['dias'] <= 30: return 'Vender con prioridad'
  if (l['reg'].year, l['reg'].month) != (cierre.year, cierre.month): return 'Recontar este mes'
  return 'Seguimiento'
estado = lambda d: 'Vencido' if d < 0 else 'Crítico ≤30' if d <= 30 else 'Alto 31–60' if d <= 60 else 'Medio 61+'

FONT = 'Arial'; H = PatternFill('solid', fgColor='D4161E')
def libro(dst):
  al = dst['alcance']
  sucs = [i for i, s in enumerate(D['s']) if al.get('todas') or s[0] == al.get('unidad') or s[1] == al.get('zona')]
  names = {D['s'][i][2]: i for i in sucs}
  L = sorted([l for l in lotes if l['suc'] in names], key=lambda l: (D['s'][names[l['suc']]][0], D['s'][names[l['suc']]][1], corto(l['suc']), l['dias'], l['prod']))
  wb = Workbook()
  # ---------- Detalle
  ws = wb.active; ws.title = 'Detalle'
  ws['A1'] = f"Vencimientos en tienda · {dst['nombre']} · cierre al {cierre:%d/%m/%Y}"; ws['A1'].font = Font(name=FONT, bold=True, size=13)
  ws['A2'] = f"Último conteo de cada producto por sucursal (últimos 2 meses), ajustado con el inventario en sistema. Los críticos (≤30 días) con conteo de hace más de 7 días se marcan para recontar. Días contados al {hoy:%d/%m/%Y}. Usa las flechas de los títulos para filtrar."
  ws['A2'].font = Font(name=FONT, size=9, italic=True, color='5A4F50')
  cols = ['Unidad', 'Zona', 'Sucursal', 'Marca', 'Código', 'Producto', 'Inv. al conteo (u)', 'Inv. sistema hoy (u)', 'Reportado (u)', 'Por vencer ajustado (u)', 'Vencimiento', 'Días para vencer', 'Estado', 'Acción', 'Colaborador', 'Fecha del conteo', 'Días desde el conteo', 'Nuevo ayer']
  ws.append([]); ws.append(cols)
  for l in L:
    s = D['s'][names[l['suc']]]
    ws.append([s[0], corto(s[1]), corto(s[2]), l['marca'], l['cod'], l['prod'], l['inv'], l['inv_hoy'], l['rep'], l['u'], l['v'], l['dias'], estado(l['dias']), accion(l), l['col'], l['reg'], l['edad'], 'Sí' if l['nuevo'] else 'No'])
  last = max(ws.max_row, 5)
  if ws.max_row == 4: ws.append(['Sin lotes reportados'] + [''] * (len(cols) - 1)); last = 5
  for r in ws.iter_rows(min_row=5, max_row=last):
    for c in r: c.font = Font(name=FONT, size=10)
    r[4].number_format = '@'; r[10].number_format = 'dd/mm/yyyy'; r[15].number_format = 'dd/mm/yyyy hh:mm'; r[9].font = Font(name=FONT, size=10, bold=True)
  t = Table(displayName='Detalle', ref=f'A4:R{last}'); t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium3', showRowStripes=True)
  ws.add_table(t)
  ws.conditional_formatting.add(f'M5:M{last}', FormulaRule(formula=['OR($M5="Vencido",$M5="Crítico ≤30")'], fill=PatternFill('solid', fgColor='F8E0E0'), font=Font(color='B42828', bold=True)))
  ws.conditional_formatting.add(f'M5:M{last}', CellIsRule(operator='equal', formula=['"Alto 31–60"'], fill=PatternFill('solid', fgColor='FBE6DC'), font=Font(color='A3431C')))
  ws.conditional_formatting.add(f'M5:M{last}', CellIsRule(operator='equal', formula=['"Medio 61+"'], fill=PatternFill('solid', fgColor='FDF1D4'), font=Font(color='8A5A00')))
  for col, w in zip('ABCDEFGHIJKLMNOPQR', [17, 10, 24, 18, 14, 40, 10, 11, 10, 12, 12, 10, 13, 34, 18, 16, 10, 9]): ws.column_dimensions[col].width = w
  ws.conditional_formatting.add(f'N5:N{last}', FormulaRule(formula=['LEFT($N5,7)="Retirar"'], font=Font(color='B42828', bold=True)))
  ws.conditional_formatting.add(f'N5:N{last}', FormulaRule(formula=['LEFT($N5,8)="Recontar"'], font=Font(color='8A5A00', bold=True)))
  ws.conditional_formatting.add(f'Q5:Q{last}', FormulaRule(formula=['AND($Q5>7,$L5<=30)'], fill=PatternFill('solid', fgColor='FDF1D4')))
  ws.freeze_panes = 'D5'
  # ---------- Resumen por sucursal
  rs = wb.create_sheet('Resumen por sucursal')
  rs['A1'] = f"Avance del mes de {mes_txt} y unidades por vencer · cierre al {cierre:%d/%m/%Y}"; rs['A1'].font = Font(name=FONT, bold=True, size=13)
  rc = ['Unidad', 'Zona', 'Sucursal', 'SKU con inventario', 'Revisados mes', '% avance', 'SKU con vencimiento', 'Vencido (u)', '≤30 días (u)', '31–60 días (u)', '61+ días (u)', 'Último registro', 'Registró']
  rs.append([]); rs.append(rc)
  tot = {'lotes': 0, 'crit': 0}
  for i in sorted(sucs, key=lambda i: (D['s'][i][0], D['s'][i][1], D['s'][i][2])):
    s = D['s'][i]; a = datos['avance'].get(f'{mes}|{s[2]}', {})
    ll = [l for l in L if l['suc'] == s[2]]; b = [sum(l['u'] for l in ll if c(l['dias'])) for c in (lambda d: d < 0, lambda d: 0 <= d <= 30, lambda d: 31 <= d <= 60, lambda d: d > 60)]
    rev = a.get('n', 0) + a.get('e', 0); r = rs.max_row + 1
    rs.append([s[0], corto(s[1]), corto(s[2]), len(D['i'][i]) // 2, rev, f'=IFERROR(E{r}/D{r},0)', len({l['cod'] for l in ll}), *b,
               dt.datetime.fromisoformat(a['act']) if a.get('act') else None, (a.get('quien') or '').split(' · ')[0]])
    tot['lotes'] += len(ll); tot['crit'] += b[0] + b[1]
  lr = rs.max_row
  for r in rs.iter_rows(min_row=4, max_row=lr):
    for c in r: c.font = Font(name=FONT, size=10)
    r[5].number_format = '0%'; r[11].number_format = 'dd/mm/yyyy hh:mm'
  rt = Table(displayName='Resumen', ref=f'A3:M{lr}'); rt.tableStyleInfo = TableStyleInfo(name='TableStyleMedium3', showRowStripes=True); rs.add_table(rt)
  rs.cell(lr + 1, 3, 'Total').font = Font(name=FONT, bold=True)
  for c, L_ in zip(range(4, 12), 'DEFGHIJK'):
    rs.cell(lr + 1, c, f'=IFERROR(E{lr+1}/D{lr+1},0)' if L_ == 'F' else f'=SUBTOTAL(109,{L_}4:{L_}{lr})').font = Font(name=FONT, bold=True)
  rs.cell(lr + 1, 6).number_format = '0%'
  for col, w in zip('ABCDEFGHIJKLM', [17, 10, 26, 12, 12, 10, 12, 11, 11, 11, 11, 16, 18]): rs.column_dimensions[col].width = w
  rs.freeze_panes = 'D4'
  wb.save(os.path.join(out, dst['archivo']))
  return {'correo': dst['correo'], 'archivo': dst['archivo'], 'sucursales': len(sucs), 'lotes': tot['lotes'], 'u_criticas': tot['crit'],
          'nuevos_ayer': sum(1 for l in L if l['nuevo'])}

res = [libro(d) for d in json.load(open(dest, encoding='utf-8'))]
print(json.dumps({'cierre': cierre.isoformat(), 'mes': mes, 'destinatarios': res}, ensure_ascii=False))
